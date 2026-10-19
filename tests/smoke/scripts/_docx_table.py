"""docx 三线表与排版辅助（对应附录 H.7 与 A.2）。"""
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _border(cell, edge: str, sz: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")
    borders.append(el)


def three_line_table(doc, df, caption: str, note: str = "", font_size: float = 10.5) -> None:
    """写标准三线表：顶线 1.5pt、栏目线 0.75pt、底线 1.5pt。"""
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(3)
    r = cap.add_run(caption)
    r.font.name = "黑体"
    r.font.size = Pt(font_size)
    r.bold = True
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    t = doc.add_table(rows=df.shape[0] + 1, cols=df.shape[1])
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for j, col in enumerate(df.columns):
        t.cell(0, j).text = str(col)
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for j, val in enumerate(row):
            t.cell(i, j).text = "" if val is None else str(val)

    for i in range(df.shape[0] + 1):
        for j in range(df.shape[1]):
            for p in t.cell(i, j).paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "宋体"
                    run.font.size = Pt(font_size)
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for j in range(df.shape[1]):
        _border(t.cell(0, j), "top", 12)
        _border(t.cell(0, j), "bottom", 6)
        _border(t.cell(df.shape[0], j), "bottom", 12)

    if note:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(note)
        r.font.name = "宋体"
        r.font.size = Pt(9)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_figure(doc, path, caption: str, width_cm: float = 13.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.font.name = "黑体"
    r.font.size = Pt(10.5)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
